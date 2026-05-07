"""Document processing for various file formats."""
import logging
from pathlib import Path
from typing import Dict, Any, Optional
import json

from .config import Config
from .file_operations import FileOperations

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Process various document formats."""

    def __init__(self, config: Config, file_ops: FileOperations):
        """Initialize document processor.

        Args:
            config: Application configuration
            file_ops: File operations handler
        """
        self.config = config
        self.file_ops = file_ops

    async def extract_data(self, file_path: Path) -> Dict[str, Any]:
        """Extract data from a document.

        Args:
            file_path: Path to document

        Returns:
            Extracted data
        """
        suffix = file_path.suffix.lower()

        handlers = {
            '.txt': self._extract_text,
            '.json': self._extract_json,
            '.csv': self._extract_csv,
            '.pdf': self._extract_pdf,
            '.docx': self._extract_docx,
            '.xlsx': self._extract_excel,
        }

        handler = handlers.get(suffix, self._extract_generic)
        return await handler(file_path)

    async def _extract_text(self, file_path: Path) -> Dict[str, Any]:
        """Extract data from text file."""
        content = self.file_ops.read_file(file_path)
        return {
            "type": "text",
            "content": content,
            "line_count": len(content.splitlines()),
            "char_count": len(content)
        }

    async def _extract_json(self, file_path: Path) -> Dict[str, Any]:
        """Extract data from JSON file."""
        content = self.file_ops.read_file(file_path)
        try:
            data = json.loads(content)
            return {
                "type": "json",
                "data": data,
                "keys": list(data.keys()) if isinstance(data, dict) else None
            }
        except json.JSONDecodeError as e:
            return {
                "type": "json",
                "error": f"Invalid JSON: {str(e)}"
            }

    async def _extract_csv(self, file_path: Path) -> Dict[str, Any]:
        """Extract data from CSV file."""
        content = self.file_ops.read_file(file_path)
        lines = content.splitlines()

        if not lines:
            return {"type": "csv", "rows": 0}

        # Simple CSV parsing (for basic cases)
        rows = [line.split(',') for line in lines]

        return {
            "type": "csv",
            "headers": rows[0] if rows else [],
            "row_count": len(rows) - 1,
            "column_count": len(rows[0]) if rows else 0,
            "preview": rows[:5]
        }

    async def _extract_pdf(self, file_path: Path) -> Dict[str, Any]:
        """Extract data from PDF file."""
        try:
            from PyPDF2 import PdfReader

            full_path = self.config.workspace_path / file_path
            reader = PdfReader(str(full_path))

            text_content = []
            for page in reader.pages:
                text_content.append(page.extract_text())

            return {
                "type": "pdf",
                "page_count": len(reader.pages),
                "content": "\n\n".join(text_content),
                "metadata": reader.metadata
            }
        except ImportError:
            return {"type": "pdf", "error": "PyPDF2 not installed"}
        except Exception as e:
            return {"type": "pdf", "error": str(e)}

    async def _extract_docx(self, file_path: Path) -> Dict[str, Any]:
        """Extract data from Word document."""
        try:
            from docx import Document

            full_path = self.config.workspace_path / file_path
            doc = Document(str(full_path))

            paragraphs = [p.text for p in doc.paragraphs]

            return {
                "type": "docx",
                "paragraph_count": len(paragraphs),
                "content": "\n".join(paragraphs),
                "tables": len(doc.tables)
            }
        except ImportError:
            return {"type": "docx", "error": "python-docx not installed"}
        except Exception as e:
            return {"type": "docx", "error": str(e)}

    async def _extract_excel(self, file_path: Path) -> Dict[str, Any]:
        """Extract data from Excel file."""
        try:
            from openpyxl import load_workbook

            full_path = self.config.workspace_path / file_path
            wb = load_workbook(str(full_path), read_only=True)

            sheets_data = {}
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheets_data[sheet_name] = {
                    "rows": sheet.max_row,
                    "columns": sheet.max_column
                }

            return {
                "type": "excel",
                "sheets": list(wb.sheetnames),
                "sheet_data": sheets_data
            }
        except ImportError:
            return {"type": "excel", "error": "openpyxl not installed"}
        except Exception as e:
            return {"type": "excel", "error": str(e)}

    async def _extract_generic(self, file_path: Path) -> Dict[str, Any]:
        """Generic extraction for unknown file types."""
        try:
            content = self.file_ops.read_file(file_path)
            return {
                "type": "generic",
                "content": content[:1000],  # First 1000 chars
                "size": len(content)
            }
        except Exception as e:
            return {
                "type": "generic",
                "error": str(e)
            }
